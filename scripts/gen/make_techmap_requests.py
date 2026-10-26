"""Techmap Request 엑셀 양식 · 예시 생성:  python scripts/gen/make_techmap_requests.py  (openpyxl 필요)
public/techmap/techmap_request_template.xlsx · techmap_request_dk_sample.xlsx · techmap_request_master_sample.xlsx
컬럼은 src/apps/techmap/lib/requests.js 의 COLUMNS와 같아야 한다."""
import pathlib
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

COLUMNS = [
    ('req_id','Req_ID',12,'고유 ID (예: REQ-2026-001)'),
    ('type','Type',16,'ADD / CHANGE_DISPLAY / RETIRE / HIDE / SHOW / OPTION_ADD / OPTION_REMOVE'),
    ('layer','Layer',14,'Layer 이름'),
    ('purpose','Purpose',12,'Purpose 이름 (RETIRE에서 비우면 Layer 전체)'),
    ('layer_no','Layer_No',10,'Master ADD · 신규 Layer 번호 (DK는 비우면 Master에서)'),
    ('abbrev','Abbrev',10,'Master ADD · 신규 Layer 약어'),
    ('purpose_no','Purpose_No',11,'Master ADD · 신규 Purpose 번호'),
    ('gds_layer','GDS_Layer',11,'Master ADD · stream 번호 (DK는 비우면 Master에서)'),
    ('gds_datatype','GDS_Datatype',13,'Master ADD · datatype'),
    ('packet','Packet',14,'기존 packet 재사용 / 공유 packet 직접 수정 시 현재 이름'),
    ('fill_color','Fill_Color',12,'색 이름 · #RRGGBB · R,G,B (DK ADD에서 비우면 Master 표시)'),
    ('outline_color','Outline_Color',14,'비우면 Fill_Color와 같게'),
    ('stipple','Stipple',11,'drDefineStipple 이름'),
    ('linestyle','LineStyle',11,'drDefineLineStyle 이름'),
    ('fill_style','Fill_Style',13,'outline / solid / X / stipple / outlineStipple …'),
    ('option','Option',14,'OPTION_ADD/REMOVE 대상 Option 이름'),
    ('requester','Requester',12,'요청자 / Site'),
    ('reason','Reason',30,'요청 사유 / Site Request 원문 요약'),
]
TYPES = 'ADD,CHANGE_DISPLAY,RETIRE,HIDE,SHOW,OPTION_ADD,OPTION_REMOVE'

def write(path, rows):
    wb = Workbook(); ws = wb.active; ws.title = 'Request'
    for c,(k,h,w,help_) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=h); cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='1F2937')
        ws.column_dimensions[cell.column_letter].width = w
        hc = ws.cell(row=2, column=c, value=help_); hc.font = Font(italic=True, color='6B7280', size=9); hc.alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(row=2, column=1, value='# 설명행(삭제 가능)').font = Font(italic=True, color='6B7280', size=9)
    ws.row_dimensions[2].height = 48; ws.freeze_panes = 'A3'
    dv = DataValidation(type='list', formula1=f'"{TYPES}"', allow_blank=True); ws.add_data_validation(dv); dv.add('B3:B500')
    for r,row in enumerate(rows, start=3):
        for c,(k,*_) in enumerate(COLUMNS, start=1):
            if row.get(k) not in (None, ''): ws.cell(row=r, column=c, value=row[k])
    g = wb.create_sheet('Guide')
    guide = [('Type','쓰는 곳','처리'),
        ('ADD','Master · DK','Master: 번호를 직접 정함(Layer_No·GDS 필수). DK: Layer·Purpose만 쓰면 번호·GDS·표시를 Master에서 가져옴, 다르게 쓰면 ERROR'),
        ('CHANGE_DISPLAY','Master · DK','packet 색·패턴 변경. 공유 packet이면 전용 packet 생성(검토 필요)'),
        ('RETIRE','Master · DK','지우지 않고 주석 처리 + RETIRED 표시. Purpose를 비우면 Layer 전체'),
        ('HIDE / SHOW','Master · DK','techDisplays Vis·Sel 플래그 nil / t (Layer는 남고 사용자 화면에서 숨김·선택 불가)'),
        ('OPTION_ADD / OPTION_REMOVE','DK','Option(Metal/Via Stack) 정의에 Layer 예외 넣기/빼기 — Option 칸 필수')]
    for r,row in enumerate(guide, start=1):
        for c,v in enumerate(row, start=1):
            cell = g.cell(row=r, column=c, value=v); cell.alignment = Alignment(wrap_text=True, vertical='top')
            if r == 1: cell.font = Font(bold=True)
    for col,w in zip('ABC',(26,14,90)): g.column_dimensions[col].width = w
    wb.save(path)

out = pathlib.Path(__file__).resolve().parents[2] / 'public/techmap'
out.mkdir(parents=True, exist_ok=True)
write(out/'techmap_request_template.xlsx', [])
write(out/'techmap_request_dk_sample.xlsx', [
    dict(req_id='REQ-101', type='ADD', layer='DMEXCL', purpose='drawing', requester='Site-A', reason='Dummy 제외 영역 사용 — 번호·표시는 Master 그대로'),
    dict(req_id='REQ-102', type='ADD', layer='AP', purpose='pin', gds_layer=60, gds_datatype=11, requester='Site-A', reason='AP pin 추가 (GDS를 잘못 적은 예)'),
    dict(req_id='REQ-103', type='ADD', layer='M11', purpose='fillDummy', requester='Site-B', reason='M11 dummy fill'),
    dict(req_id='REQ-104', type='HIDE', layer='text', purpose='drawing', requester='Site-B', reason='text layer 사용자 선택 막기'),
    dict(req_id='REQ-105', type='SHOW', layer='prBoundary', purpose='drawing', requester='Site-B', reason='prBoundary 다시 보이게'),
    dict(req_id='REQ-106', type='OPTION_ADD', layer='PAD', option='1P9M_6X1Z', requester='Site-C', reason='9M Option에도 PAD 제공'),
    dict(req_id='REQ-107', type='OPTION_REMOVE', layer='VIA10', option='1P11M_8X2Z', requester='Site-C', reason='VIA10 제외 요청 (Via 스택이 끊기는 예)'),
    dict(req_id='REQ-108', type='CHANGE_DISPLAY', layer='M5', purpose='drawing', fill_color='#3366FF', requester='Site-C', reason='M5 색 구분'),
    dict(req_id='REQ-109', type='ADD', layer='NEWX', purpose='drawing', requester='Site-D', reason='Master에 없는 Layer 예'),
])
write(out/'techmap_request_master_sample.xlsx', [
    dict(req_id='MREQ-01', type='ADD', layer='M12', purpose='drawing', layer_no=33, abbrev='M12', gds_layer=33, gds_datatype=0, fill_color='olive', stipple='slash', linestyle='solid', fill_style='outlineStipple', requester='공정팀', reason='12M 확장'),
    dict(req_id='MREQ-02', type='ADD', layer='VIA11', purpose='drawing', layer_no=32, abbrev='V11', gds_layer=32, gds_datatype=0, fill_color='teal', stipple='cross', linestyle='solid', fill_style='outlineStipple', requester='공정팀', reason='12M 확장'),
    dict(req_id='MREQ-03', type='ADD', layer='MIM', purpose='drawing', layer_no=70, abbrev='MIM', gds_layer=71, gds_datatype=0, fill_color='pink', requester='공정팀', reason='SF2와 GDS가 다른 예 (SF2 MIM = 70/0)'),
])
print('written', out)
